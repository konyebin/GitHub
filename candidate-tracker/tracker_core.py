"""Parsing, scoring and Excel export for the candidate tracker (no UI code)."""

from __future__ import annotations

import colorsys
import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import get_column_letter

HIRE, STANDBY, NO_HIRE, UNASSIGNED = "Hire", "Standby", "No Hire", "Unassigned"
STATUSES = [HIRE, STANDBY, NO_HIRE]
ALL_STATUSES = [UNASSIGNED, HIRE, STANDBY, NO_HIRE]

# UI colours (dark theme) and Excel colours (fill, font) per status.
UI_COLORS = {
    HIRE: "#2ECC71",
    STANDBY: "#F5C542",
    NO_HIRE: "#FF5C5C",
    UNASSIGNED: "#7C8BA1",
}
XL_COLORS = {
    HIRE: ("C6EFCE", "006100"),
    STANDBY: ("FFEB9C", "7F6000"),
    NO_HIRE: ("FFC7CE", "9C0006"),
    UNASSIGNED: ("EDEFF3", "4A5568"),
}
CISCO_BLUE = "049FD9"
NAVY = "0A1220"

ITEM_MAX = 4
SECTION_MAX = 16
TOTAL_MAX = 48

SECTIONS = {
    "interview": "Interview",
    "role play": "Role Play",
    "presentation": "Presentation",
}

DEFAULT_HIRE_THRESHOLD = 0.85
DEFAULT_STANDBY_THRESHOLD = 0.70


@dataclass
class ParseResult:
    df: pd.DataFrame
    skipped_examples: int


def _num(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        text = str(value).strip()
        return float(text) if text else None
    except ValueError:
        return None


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def normalise_status(value: Any) -> str:
    """Map free-text status cells ('Hire', 'no hire', 'Stand by', ...) to a tracker status."""
    text = re.sub(r"[^a-z]", "", _text(value).lower())
    if not text:
        return UNASSIGNED
    if "standby" in text or text in {"hold", "waitlist", "waitlisted"}:
        return STANDBY
    if "nohire" in text or text in {"reject", "rejected", "decline", "declined", "no"}:
        return NO_HIRE
    if "hire" in text or text in {"yes", "offer", "selected"}:
        return HIRE
    return UNASSIGNED


def _fill_status(cell) -> str:
    """Infer a status from a solid green / yellow / red cell fill, if there is one."""
    fill = cell.fill
    if fill is None or fill.fill_type != "solid":
        return UNASSIGNED
    color = fill.fgColor
    rgb = None
    if color.type == "rgb" and isinstance(color.rgb, str):
        rgb = color.rgb[-6:]
    elif color.type == "indexed" and isinstance(color.indexed, int) and color.indexed < len(COLOR_INDEX):
        rgb = COLOR_INDEX[color.indexed][-6:]
    if not rgb or not re.fullmatch(r"[0-9A-Fa-f]{6}", rgb):
        return UNASSIGNED
    r, g, b = (int(rgb[i : i + 2], 16) / 255 for i in (0, 2, 4))
    hue, sat, val = colorsys.rgb_to_hsv(r, g, b)
    if sat < 0.12 or val < 0.35:
        return UNASSIGNED
    degrees = hue * 360
    if degrees <= 20 or degrees >= 340:
        return NO_HIRE
    if 20 < degrees < 75:
        return STANDBY
    if 75 <= degrees <= 170:
        return HIRE
    return UNASSIGNED


def _find_header_row(ws) -> int:
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 40)):
        for cell in row:
            if _text(cell.value).lower() == "candidate name":
                return cell.row
    raise ValueError(
        "Could not find a 'Candidate Name' header row. "
        "Is this the FY27Q3 ASE Scores workbook?"
    )


def _section_columns(headers: list[str]) -> dict[str, dict[str, int]]:
    """Locate each scoring section from its 'TOTAL - <section>' column.

    The four competency columns sit immediately left of the total, and the
    recommendation / comments columns immediately right of it.
    """
    found: dict[str, dict[str, int]] = {}
    for idx, header in enumerate(headers):
        low = header.lower()
        if not low.startswith("total"):
            continue
        for needle, label in SECTIONS.items():
            if needle in low and label not in found:
                found[label] = {
                    "total": idx,
                    "items": list(range(max(idx - 4, 0), idx)),
                    "rec": idx + 1,
                }
    missing = [label for label in SECTIONS.values() if label not in found]
    if missing:
        raise ValueError(f"Missing score sections in workbook: {', '.join(missing)}")
    return found


def _section_score(item_values: list[float | None], cached_total: float | None) -> float | None:
    present = [v for v in item_values if v is not None]
    if present:
        return round(sum(present) / (len(present) * ITEM_MAX) * SECTION_MAX, 2)
    if cached_total:
        return round(cached_total, 2)
    return None


def parse_workbook(source: Any, include_examples: bool = False) -> ParseResult:
    """Read candidates from the ASE scores workbook (path, file object or bytes)."""
    if isinstance(source, (bytes, bytearray)):
        source = io.BytesIO(source)
    wb = openpyxl.load_workbook(source, data_only=True)
    ws = wb.worksheets[0]
    header_row = _find_header_row(ws)
    width = ws.max_column
    headers = [_text(ws.cell(header_row, c).value) for c in range(1, width + 1)]
    lowered = [h.lower() for h in headers]

    def col(name: str, default: int | None = None) -> int | None:
        return lowered.index(name) if name in lowered else default

    name_col = col("candidate name")
    country_col = col("country")
    salutation_col = col("salutation")
    buddy_col = col("buddy")
    sections = _section_columns(headers)
    status_col = next(
        (
            i
            for i, h in enumerate(lowered)
            if re.sub(r"\s+", " ", h).startswith("hire") and "standby" in h
        ),
        None,
    )
    notes_col = None
    if buddy_col is not None:
        for idx in range(buddy_col + 1, width):
            if lowered[idx].startswith("comments"):
                notes_col = idx
                break

    rows: list[dict[str, Any]] = []
    seen: dict[str, int] = {}
    skipped = 0
    for r in range(header_row + 1, ws.max_row + 1):
        values = [ws.cell(r, c).value for c in range(1, width + 1)]
        name = _text(values[name_col])
        if not name:
            continue
        if re.match(r"^example\b", name, re.IGNORECASE) and not include_examples:
            skipped += 1
            continue
        seen[name.lower()] = seen.get(name.lower(), 0) + 1
        if seen[name.lower()] > 1:
            name = f"{name} ({seen[name.lower()]})"

        record: dict[str, Any] = {
            "Candidate": name,
            "Country": _text(values[country_col]) if country_col is not None else "",
            "Salutation": _text(values[salutation_col]) if salutation_col is not None else "",
        }
        total = 0.0
        scored = 0
        for label, cols in sections.items():
            items = [_num(values[i]) for i in cols["items"]]
            score = _section_score(items, _num(values[cols["total"]]))
            record[label] = score
            record[f"{label} Rec"] = _text(values[cols["rec"]]) if cols["rec"] < width else ""
            for n, item in enumerate(items, start=1):
                record[f"{label} C{n}"] = item
            if score is not None:
                total += score
                scored += 1
        record["Total"] = round(total, 2)
        record["Overall %"] = round(total / TOTAL_MAX, 4)
        record["Complete"] = scored == len(sections)
        source_status = UNASSIGNED
        if status_col is not None:
            source_status = normalise_status(values[status_col])
            if source_status == UNASSIGNED:
                source_status = _fill_status(ws.cell(r, status_col + 1))
        if source_status == UNASSIGNED:
            source_status = _fill_status(ws.cell(r, name_col + 1))
        record["Source Status"] = source_status
        record["Buddy"] = _text(values[buddy_col]) if buddy_col is not None else ""
        record["Source Notes"] = _text(values[notes_col]) if notes_col is not None else ""
        rows.append(record)

    df = pd.DataFrame(rows)
    return ParseResult(df=df, skipped_examples=skipped)


def suggest(pct: float, complete: bool, hire_th: float, standby_th: float) -> str:
    if not complete:
        return "Incomplete"
    if pct >= hire_th:
        return HIRE
    if pct >= standby_th:
        return STANDBY
    return NO_HIRE


def normalise_emp(value: Any) -> str:
    return _text(value)


def duplicate_numbers(emp_numbers: dict[str, str]) -> dict[str, list[str]]:
    """Map each duplicated employee number (case-insensitive) to its candidates."""
    groups: dict[str, list[str]] = {}
    for candidate, number in emp_numbers.items():
        key = normalise_emp(number).lower()
        if key:
            groups.setdefault(key, []).append(candidate)
    return {k: v for k, v in groups.items() if len(v) > 1}


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", start_color=hex_color, end_color=hex_color)


_THIN = Side(style="thin", color="D5DAE1")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _header(ws, row: int, labels: list[str]) -> None:
    for c, label in enumerate(labels, start=1):
        cell = ws.cell(row, c, label)
        cell.fill = _fill(NAVY)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="thick", color=CISCO_BLUE))
    ws.row_dimensions[row].height = 30


def _style_status_cell(cell, status: str, bold: bool = False) -> None:
    fill, font = XL_COLORS[status]
    cell.fill = _fill(fill)
    cell.font = Font(color=font, bold=bold)
    cell.border = _BORDER


def _autowidth(ws, minimum: int = 10, maximum: int = 48) -> None:
    for column in ws.columns:
        letter = get_column_letter(column[0].column)
        longest = max((len(str(c.value)) for c in column if c.value is not None), default=0)
        ws.column_dimensions[letter].width = max(minimum, min(maximum, longest + 3))


def _as_date_text(value: Any) -> str:
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    return _text(value)


def build_workbook(
    candidates: pd.DataFrame,
    audit: list[dict[str, Any]],
    hire_th: float,
    standby_th: float,
) -> bytes:
    """Create the colour-coded tracker workbook and return it as bytes.

    ``candidates`` needs the parsed score columns plus Status, Employee #,
    Re-evaluate On and Decision Notes.
    """
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "All Candidates"
    cols = [
        "Candidate", "Country", "Interview", "Role Play", "Presentation", "Total /48",
        "Overall %", "Status", "Employee #", "Re-evaluate On", "Buddy", "Notes",
    ]
    _header(ws, 1, cols)
    for r, row in enumerate(candidates.to_dict("records"), start=2):
        status = row["Status"]
        notes = " | ".join(n for n in (_text(row["Decision Notes"]), _text(row["Source Notes"])) if n)
        values = [
            row["Candidate"], row["Country"], row["Interview"], row["Role Play"],
            row["Presentation"], row["Total"], row["Overall %"],
            status, _text(row["Employee #"]), _as_date_text(row["Re-evaluate On"]),
            row["Buddy"], notes,
        ]
        for c, value in enumerate(values, start=1):
            cell = ws.cell(r, c, None if (isinstance(value, float) and pd.isna(value)) else value)
            _style_status_cell(cell, status, bold=(c in (1, 8)))
            if c == 7:
                cell.number_format = "0.0%"
            elif c in (3, 4, 5, 6):
                cell.number_format = "0.0"
            if c in (3, 4, 5, 6, 7, 8):
                cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(ws.max_row, 2)}"
    _autowidth(ws)

    board = wb.create_sheet("By Status")
    names = {s: candidates[candidates["Status"] == s] for s in ALL_STATUSES}
    order = [HIRE, STANDBY, NO_HIRE, UNASSIGNED]
    for c, status in enumerate(order, start=1):
        head = board.cell(1, c, f"{status} ({len(names[status])})")
        fill, font = XL_COLORS[status]
        head.fill = _fill(fill)
        head.font = Font(bold=True, color=font, size=12)
        head.alignment = Alignment(horizontal="center", vertical="center")
        head.border = Border(bottom=Side(style="thick", color=CISCO_BLUE))
        for r, row in enumerate(names[status].to_dict("records"), start=2):
            label = f"{row['Candidate']}  ({row['Overall %']:.0%})"
            if _text(row["Employee #"]) and status == HIRE:
                label += f"  #{_text(row['Employee #'])}"
            cell = board.cell(r, c, label)
            _style_status_cell(cell, status)
        board.column_dimensions[get_column_letter(c)].width = 38
    board.row_dimensions[1].height = 28

    emp = wb.create_sheet("Employee Numbers")
    _header(emp, 1, ["Candidate", "Status", "Employee #", "Flag"])
    holders = candidates[
        (candidates["Status"] == HIRE) | (candidates["Employee #"].astype(str).str.strip() != "")
    ].copy()
    holders["_order"] = (holders["Status"] != HIRE).astype(int)
    holders = holders.sort_values(["_order", "Candidate"])
    dupes = duplicate_numbers(dict(zip(candidates["Candidate"], candidates["Employee #"])))
    dupe_keys = set(dupes)
    for r, row in enumerate(holders.to_dict("records"), start=2):
        number = _text(row["Employee #"])
        flag = ""
        if number.lower() in dupe_keys:
            flag = "DUPLICATE"
        elif row["Status"] == HIRE and not number:
            flag = "MISSING"
        elif row["Status"] != HIRE:
            flag = "Kept from earlier Hire"
        for c, value in enumerate([row["Candidate"], row["Status"], number, flag], start=1):
            cell = emp.cell(r, c, value)
            _style_status_cell(cell, row["Status"], bold=(c == 1))
        if flag in ("DUPLICATE", "MISSING"):
            emp.cell(r, 4).font = Font(bold=True, color="C00000")
    emp.freeze_panes = "A2"
    _autowidth(emp)

    log = wb.create_sheet("Audit Log")
    _header(log, 1, ["When", "Reviewer", "Candidate", "Change", "From", "To"])
    for r, entry in enumerate(reversed(audit), start=2):
        for c, key in enumerate(["when", "by", "candidate", "field", "old", "new"], start=1):
            cell = log.cell(r, c, _text(entry.get(key, "")))
            cell.border = _BORDER
    log.freeze_panes = "A2"
    _autowidth(log)

    summary = wb.create_sheet("Summary")
    _header(summary, 1, ["Status", "Candidates", "Avg Overall %", "Top Score %"])
    for r, status in enumerate(ALL_STATUSES, start=2):
        subset = candidates[candidates["Status"] == status]
        values = [
            status,
            len(subset),
            float(subset["Overall %"].mean()) if len(subset) else None,
            float(subset["Overall %"].max()) if len(subset) else None,
        ]
        for c, value in enumerate(values, start=1):
            cell = summary.cell(r, c, value)
            _style_status_cell(cell, status, bold=(c == 1))
            if c in (3, 4):
                cell.number_format = "0.0%"
    summary.cell(7, 1, "Total candidates").font = Font(bold=True)
    summary.cell(7, 2, len(candidates)).font = Font(bold=True)
    summary.cell(9, 1, "Suggestion thresholds").font = Font(bold=True, color=CISCO_BLUE)
    summary.cell(10, 1, "Hire at or above")
    summary.cell(10, 2, hire_th).number_format = "0%"
    summary.cell(11, 1, "Standby at or above")
    summary.cell(11, 2, standby_th).number_format = "0%"
    summary.cell(12, 1, "Generated")
    summary.cell(12, 2, datetime.now().strftime("%Y-%m-%d %H:%M"))
    _autowidth(summary)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
