from pathlib import Path

import openpyxl
import pandas as pd

from tracker_core import add_signals, HIRE, NO_HIRE, STANDBY, UNASSIGNED, build_workbook, duplicate_numbers, normalise_status, parse_workbook, suggest

SAMPLES = Path(__file__).resolve().parent.parent / "samples"


def test_template_has_only_example_row():
    res = parse_workbook(SAMPLES / "FY27Q3_AMER_ASE_Scores_template.xlsx")
    assert res.df.empty and res.skipped_examples == 1
    with_examples = parse_workbook(SAMPLES / "FY27Q3_AMER_ASE_Scores_template.xlsx", include_examples=True)
    row = with_examples.df.iloc[0]
    assert row["Total"] == 48 and row["Overall %"] == 1.0


def test_demo_parse_and_scores():
    df = parse_workbook(SAMPLES / "FY27Q3_AMER_ASE_Scores_TEST_DATA.xlsx").df
    assert len(df) == 20
    assert (df["Interview"] <= 16).all()
    assert not df.iloc[-1]["Complete"]
    assert df["Complete"].iloc[:-1].all()


def test_suggest_thresholds():
    assert suggest(0.9, True, 0.85, 0.7) == HIRE
    assert suggest(0.75, True, 0.85, 0.7) == STANDBY
    assert suggest(0.5, True, 0.85, 0.7) == NO_HIRE
    assert suggest(0.9, False, 0.85, 0.7) == "Incomplete"


def test_duplicates_case_insensitive():
    assert duplicate_numbers({"a": "E1", "b": "e1", "c": "", "d": ""}) == {"e1": ["a", "b"]}


def test_export_colours_and_sheets(tmp_path):
    df = parse_workbook(SAMPLES / "FY27Q3_AMER_ASE_Scores_TEST_DATA.xlsx").df
    df["Status"] = [HIRE, STANDBY, NO_HIRE] + [UNASSIGNED] * (len(df) - 3)
    df["Employee #"] = ["E100", "E100", ""] + [""] * (len(df) - 3)
    df["Re-evaluate On"] = None
    df["Decision Notes"] = ""
    data = build_workbook(df, [{"when": "now", "by": "me", "candidate": "x", "field": "Status", "old": "-", "new": "Hire"}], 0.85, 0.7)
    path = tmp_path / "out.xlsx"
    path.write_bytes(data)
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["All Candidates", "By Status", "Needs Review", "Employee Numbers", "Audit Log", "Summary"]
    ws = wb["All Candidates"]
    assert ws["A2"].fill.start_color.rgb.endswith("C6EFCE")
    assert ws["A3"].fill.start_color.rgb.endswith("FFEB9C")
    assert ws["A4"].fill.start_color.rgb.endswith("FFC7CE")
    emp = wb["Employee Numbers"]
    flags = [emp.cell(r, 4).value for r in range(2, emp.max_row + 1)]
    assert flags.count("DUPLICATE") == 2


def test_source_status_is_imported():
    df = parse_workbook(SAMPLES / "FY27Q3_AMER_ASE_Scores_TEST_DATA.xlsx").df
    counts = df["Source Status"].value_counts().to_dict()
    assert counts.get(HIRE, 0) > 0 and counts.get(STANDBY, 0) > 0 and counts.get(NO_HIRE, 0) > 0
    assert counts.get(UNASSIGNED, 0) == 3


def test_normalise_status_variants():
    assert normalise_status("Hire") == HIRE
    assert normalise_status(" no hire ") == NO_HIRE
    assert normalise_status("No-Hire") == NO_HIRE
    assert normalise_status("Stand by") == STANDBY
    assert normalise_status("STANDBY") == STANDBY
    assert normalise_status(None) == UNASSIGNED


def test_status_from_cell_fill(tmp_path):
    from openpyxl.styles import PatternFill

    src = SAMPLES / "FY27Q3_AMER_ASE_Scores_TEST_DATA.xlsx"
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    for r in range(11, 31):
        ws.cell(r, 27).value = None
    for r, color in ((11, "C6EFCE"), (12, "FFEB9C"), (13, "FFC7CE"), (14, "92D050"), (15, "FF0000")):
        ws.cell(r, 27).fill = PatternFill("solid", start_color=color, end_color=color)
    path = tmp_path / "filled.xlsx"
    wb.save(path)
    df = parse_workbook(path).df
    assert list(df["Source Status"].iloc[:5]) == [HIRE, STANDBY, NO_HIRE, HIRE, NO_HIRE]
    assert (df["Source Status"].iloc[5:] == UNASSIGNED).all()


def test_score_band_boundaries():
    assert suggest(0.85, True, 0.85, 0.70) == HIRE
    assert suggest(0.8499, True, 0.85, 0.70) == STANDBY
    assert suggest(0.70, True, 0.85, 0.70) == STANDBY
    assert suggest(0.6999, True, 0.85, 0.70) == NO_HIRE


def test_anomalies():
    from tracker_core import detect_anomaly

    assert detect_anomaly(HIRE, 0.60, True, 0.85, 0.70)[0] == "high"
    assert detect_anomaly(HIRE, 0.78, True, 0.85, 0.70)[0] == "review"
    assert detect_anomaly(HIRE, 0.90, True, 0.85, 0.70)[0] == ""
    assert detect_anomaly(NO_HIRE, 0.92, True, 0.85, 0.70)[0] == "high"
    assert detect_anomaly(NO_HIRE, 0.75, True, 0.85, 0.70)[0] == "review"
    assert detect_anomaly(NO_HIRE, 0.40, True, 0.85, 0.70)[0] == ""
    assert detect_anomaly(STANDBY, 0.78, True, 0.85, 0.70)[0] == ""
    assert detect_anomaly(STANDBY, 0.95, True, 0.85, 0.70)[0] == "review"
    assert detect_anomaly(UNASSIGNED, 0.10, True, 0.85, 0.70)[0] == ""
    assert detect_anomaly(HIRE, 0.90, False, 0.85, 0.70)[0] == "review"
    assert "note added" in detect_anomaly(HIRE, 0.60, True, 0.85, 0.70, noted=True)[1]


def test_add_signals_columns():
    df = parse_workbook(SAMPLES / "FY27Q3_AMER_ASE_Scores_TEST_DATA.xlsx").df
    df["Status"] = HIRE
    df["Decision Notes"] = ""
    out = add_signals(df, 0.85, 0.70)
    assert set(out["Score Band"]) <= {"Hire", "Maybe", "No Hire", "Incomplete"}
    low = out[(out["Overall %"] < 0.70) & out["Complete"]]
    assert (low["Flag Level"] == "high").all() and low["Flag"].str.startswith("HIGH").all()


def test_template_without_candidates_returns_typed_empty_frame():
    res = parse_workbook(SAMPLES / "FY27Q3_AMER_ASE_Scores_template.xlsx")
    assert res.df.empty
    assert "Candidate" in res.df.columns and "Source Status" in res.df.columns
