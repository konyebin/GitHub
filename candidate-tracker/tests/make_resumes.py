"""Generate 10 fictional test resumes (.docx) for candidates in the TEST_DATA workbook.

Run:  python tests/make_resumes.py
Output: samples/resumes/<First>_<Last>_Resume.docx

Everyone here is fictional. Resume strength loosely follows each candidate's score in the test workbook so the
files are useful for demos (stronger scores get richer, more relevant experience).
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT_DIR = Path(__file__).resolve().parent.parent / "samples" / "resumes"
BLUE = RGBColor(0x04, 0x9F, 0xD9)
GREY = RGBColor(0x55, 0x5F, 0x6D)

RESUMES = [
    {
        "name": "Taylor Nunes",
        "location": "Sao Paulo, Brazil",
        "phone": "+55 11 5550 0142",
        "summary": (
            "Network engineer and aspiring Associate Systems Engineer with two internships in carrier networking and "
            "pre-sales. Trilingual (Portuguese, Spanish, English) and comfortable explaining routing, security and "
            "collaboration concepts to non-technical audiences."
        ),
        "experience": [
            ("Systems Engineering Intern", "Aurora Networks", "Sao Paulo", "Jan 2025 - Jul 2025", [
                "Built and demoed a Webex Calling and Meeting proof of concept for 6 enterprise prospects, supporting 2 closed deals.",
                "Wrote Python scripts that cut lab provisioning time from 3 hours to 20 minutes.",
                "Presented weekly technical readouts to account teams across Latin America.",
            ]),
            ("Network Operations Intern", "Meridian Telecom", "Sao Paulo", "Jun 2024 - Dec 2024", [
                "Monitored a 400-node MPLS backbone and resolved Tier 1 incidents with a 95% first-pass fix rate.",
                "Documented 15 troubleshooting runbooks adopted by the on-call team.",
            ]),
        ],
        "education": [("M.S. Computer Engineering", "Universidade Paulista do Sul", "2025", "GPA 3.9/4.0")],
        "certs": ["CCNP Enterprise (ENCOR)", "Cisco DevNet Associate", "CCNA"],
        "skills": ["Routing and switching (OSPF, BGP, VLANs)", "Python, REST APIs, Postman", "Webex Calling and Meetings",
                   "SD-WAN concepts", "Customer demos and workshops", "Portuguese, Spanish, English"],
        "projects": ["Open-source toolkit that auto-generates network diagrams from device configs (320 GitHub stars)."],
    },
    {
        "name": "Reese Duarte",
        "location": "Bogota, Colombia",
        "phone": "+57 1 555 0178",
        "summary": (
            "Telecommunications engineer with two years in a network operations center and a growing pre-sales "
            "focus. Known for calm incident handling and clear customer communication in Spanish and English."
        ),
        "experience": [
            ("Network Operations Engineer", "Andes Connect", "Bogota", "Mar 2023 - Present", [
                "Own escalations for 80+ enterprise voice and data customers; average resolution time down 28% year over year.",
                "Led migration of 1,200 users from legacy PBX to cloud calling with zero unplanned downtime.",
                "Trained 10 new hires on ticket triage and customer-facing communication.",
            ]),
            ("Technical Support Associate", "Condor Wireless", "Bogota", "Jun 2021 - Feb 2023", [
                "Resolved 35+ tickets per day with a 4.8/5 satisfaction score.",
            ]),
        ],
        "education": [("B.S. Telecommunications Engineering", "Universidad del Valle Central", "2021", "")],
        "certs": ["CCNA", "Cisco Webex Calling Specialist", "ITIL 4 Foundation"],
        "skills": ["Cloud calling and SIP troubleshooting", "Wireshark, Splunk", "Cisco switching and wireless",
                   "Incident management", "Spanish (native), English (fluent)"],
        "projects": ["Built a call-quality dashboard that flagged packet-loss hotspots across 14 customer sites."],
    },
    {
        "name": "Jordan Blake",
        "location": "Austin, TX",
        "phone": "(555) 010-2231",
        "summary": (
            "Customer-focused technical support engineer moving into pre-sales. Strong storyteller who turns complex "
            "troubleshooting into clear guidance, with hands-on collaboration and networking experience."
        ),
        "experience": [
            ("Technical Support Engineer", "BrightPath Software", "Austin, TX", "Aug 2024 - Present", [
                "Resolved complex cloud-communications cases for enterprise accounts; top 10% of 60 engineers by CSAT.",
                "Created 25 knowledge-base articles that deflected an estimated 400 tickets per quarter.",
                "Delivered 8 customer onboarding webinars to groups of 30 to 120 attendees.",
            ]),
            ("IT Intern", "Capital City Schools", "Austin, TX", "May 2023 - Aug 2023", [
                "Deployed 150 Wi-Fi access points and documented the rollout for district IT.",
            ]),
        ],
        "education": [("B.S. Computer Science", "University of Texas Hill Country", "2024", "Dean's List")],
        "certs": ["CCNA", "Cisco DevNet Associate"],
        "skills": ["Troubleshooting and root-cause analysis", "Python, Bash", "Webex, Zoom, Teams administration",
                   "Public speaking and demos", "Salesforce, Jira"],
        "projects": ["Meetup talk: 'Explaining VoIP quality to non-engineers' (120 attendees)."],
    },
    {
        "name": "Morgan Ellis",
        "location": "Toronto, ON, Canada",
        "phone": "(416) 555-0113",
        "summary": (
            "Network analyst with co-op and full-time experience in managed services. Dependable, curious and "
            "comfortable supporting live customer environments."
        ),
        "experience": [
            ("Network Analyst", "Northfield Managed Services", "Toronto, ON", "Sep 2023 - Present", [
                "Administer Meraki and Catalyst networks for 40 small and mid-size customers.",
                "Cut recurring Wi-Fi tickets 35% by standardizing SSID and RF profiles.",
                "Present quarterly business reviews to customer IT leads.",
            ]),
            ("Co-op Student, Service Desk", "Lakeshore Bank", "Toronto, ON", "May 2022 - Aug 2023", [
                "Supported 900 branch and head-office users; maintained a 96% SLA attainment.",
            ]),
        ],
        "education": [("B.Tech. Information Technology", "Ontario Polytechnic Institute", "2023", "Co-op program")],
        "certs": ["CCNA", "Meraki CMNA"],
        "skills": ["Meraki dashboard", "VLANs, STP, OSPF", "Firewall policy basics", "Client presentations",
                   "English (native), French (conversational)"],
        "projects": ["Home lab with 3 switches and 2 firewalls used to practice failover scenarios."],
    },
    {
        "name": "Drew Patel",
        "location": "Mississauga, ON, Canada",
        "phone": "(905) 555-0187",
        "summary": (
            "Recent engineering graduate with ISP internship experience, Linux and Python skills, and a steady focus "
            "on building customer-facing technical depth."
        ),
        "experience": [
            ("Network Engineering Intern", "Maple Broadband", "Mississauga, ON", "May 2024 - Aug 2024", [
                "Automated nightly config backups for 220 routers using Python and Netmiko.",
                "Shadowed field teams on 12 enterprise fibre turn-ups.",
            ]),
            ("Teaching Assistant, Computer Networks", "Greater Toronto University", "Toronto, ON", "Sep 2023 - Apr 2024", [
                "Ran weekly labs for 45 students; held office hours on subnetting and routing.",
            ]),
        ],
        "education": [("B.Eng. Computer Engineering", "Greater Toronto University", "2024", "")],
        "certs": ["CCNA (in progress, exam scheduled)"],
        "skills": ["Linux, Python, Git", "Subnetting and routing", "Wireshark", "Technical writing"],
        "projects": ["Capstone: low-cost IoT sensor network with a Grafana dashboard for campus energy use."],
    },
    {
        "name": "Quinn Harper",
        "location": "Denver, CO",
        "phone": "(555) 010-3340",
        "summary": (
            "Sales operations professional with a networking bootcamp certificate and a talent for presenting. "
            "Seeking to combine customer rapport with deeper technical skills in a systems engineering role."
        ),
        "experience": [
            ("Sales Operations Analyst", "Cobalt Cloud", "Denver, CO", "Jan 2023 - Present", [
                "Supported a 25-person sales team with pipeline reporting, quoting and deal desk workflows.",
                "Sat in on 100+ customer calls and wrote technical follow-ups for account executives.",
                "Won company 'Best Presenter' award at the annual kickoff.",
            ]),
            ("Customer Success Associate", "Summit Analytics", "Denver, CO", "Jun 2021 - Dec 2022", [
                "Managed onboarding for 60 mid-market accounts with 92% renewal.",
            ]),
        ],
        "education": [
            ("B.A. Business Administration", "Colorado Mountain University", "2021", ""),
            ("Network Engineering Bootcamp", "Rocky Mountain Tech Academy", "2023", "12 weeks, evenings"),
        ],
        "certs": ["CompTIA Network+", "CCNA Learning Path (completed coursework)"],
        "skills": ["Executive presentations", "Salesforce, Tableau", "Networking fundamentals", "Discovery and demos"],
        "projects": ["Built a customer-facing ROI calculator adopted by 6 account teams."],
    },
    {
        "name": "Alex Rivera",
        "location": "Buenos Aires, Argentina",
        "phone": "+54 11 5550 0166",
        "summary": (
            "Systems administrator with three years of hands-on infrastructure experience. Deep Linux and network "
            "troubleshooting skills; looking to move into customer-facing technical work."
        ),
        "experience": [
            ("Systems Administrator", "Pampa Logistics", "Buenos Aires", "Feb 2022 - Present", [
                "Maintain 60 Linux servers and a 12-site WAN with 99.95% availability.",
                "Introduced Ansible playbooks that cut patching effort by 50%.",
                "Led a firewall refresh across 12 branches over six weekends.",
            ]),
            ("IT Support Technician", "Rio de la Plata Insurance", "Buenos Aires", "Mar 2020 - Jan 2022", [
                "First-line support for 300 employees across two offices.",
            ]),
        ],
        "education": [("B.S. Systems Engineering", "Universidad Tecnologica del Sur", "2020", "")],
        "certs": ["CompTIA Network+", "Linux Foundation Certified Sysadmin", "CCNA (in progress)"],
        "skills": ["Linux, Ansible, Bash", "VPN and firewall management", "Zabbix, Grafana", "Spanish (native), English (B2)"],
        "projects": ["Wrote an internal guide to troubleshooting WAN latency; used in onboarding for new staff."],
    },
    {
        "name": "Riley Chen",
        "location": "Seattle, WA",
        "phone": "(555) 010-4472",
        "summary": (
            "Help desk and customer support professional with a degree in information systems. Patient communicator "
            "who enjoys learning new tools quickly."
        ),
        "experience": [
            ("IT Help Desk Analyst", "Harborlight Health", "Seattle, WA", "Jul 2024 - Present", [
                "Support 1,100 clinical and administrative users; 97% of tickets closed within SLA.",
                "Rolled out softphone clients to 200 staff with step-by-step training sessions.",
            ]),
            ("Customer Support Representative", "Evergreen Outfitters", "Seattle, WA", "Jun 2022 - Jun 2024", [
                "Handled chat and phone support for 40 customers per day; 4.7/5 satisfaction.",
            ]),
        ],
        "education": [("B.S. Information Systems", "Puget Sound State University", "2024", "")],
        "certs": ["CompTIA A+", "CompTIA Network+"],
        "skills": ["Active Directory, Microsoft 365", "Ticketing (ServiceNow)", "Basic networking", "Customer communication"],
        "projects": [],
    },
    {
        "name": "Emerson Reid",
        "location": "Atlanta, GA",
        "phone": "(555) 010-5589",
        "summary": (
            "Customer service and retail management professional exploring a move into technology. Completed "
            "introductory networking coursework and building hands-on skills."
        ),
        "experience": [
            ("Assistant Store Manager", "Peachtree Home Supply", "Atlanta, GA", "Mar 2022 - Present", [
                "Supervise a team of 14 and handle escalated customer issues.",
                "Introduced a new point-of-sale training routine that shortened onboarding by a week.",
            ]),
            ("Sales Associate", "Peachtree Home Supply", "Atlanta, GA", "Jun 2020 - Feb 2022", [
                "Consistently met monthly sales targets.",
            ]),
        ],
        "education": [("B.S. Business Administration", "Georgia Piedmont College", "2020", "")],
        "certs": ["Cisco Networking Academy: CCNA Introduction to Networks (completed)"],
        "skills": ["Team leadership", "Customer service", "Microsoft Office", "Introductory networking"],
        "projects": ["Set up a small home network lab to practice IP addressing."],
    },
    {
        "name": "Sam Ortega",
        "location": "Mexico City, Mexico",
        "phone": "+52 55 5550 0129",
        "summary": (
            "IT support technician with one year of experience and a technical diploma. Eager to grow into a "
            "networking and collaboration role."
        ),
        "experience": [
            ("IT Support Technician", "Servicios Azteca", "Mexico City", "Feb 2025 - Present", [
                "Provide desktop and printer support for 150 users in two offices.",
                "Configured basic VLANs on small-office switches under supervision.",
            ]),
            ("Computer Lab Assistant", "Colegio Tecnico Metropolitano", "Mexico City", "Aug 2023 - Dec 2024", [
                "Maintained 40 lab computers and helped students with software installs.",
            ]),
        ],
        "education": [("Technical Diploma, Computer Systems", "Colegio Tecnico Metropolitano", "2024", "")],
        "certs": ["CompTIA A+ (in progress)"],
        "skills": ["Windows support", "Basic LAN setup", "Spanish (native), English (intermediate)"],
        "projects": [],
    },
]


def _bottom_border(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "049FD9")
    borders.append(bottom)
    p_pr.append(borders)


def _heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(10.5)
    run.font.color.rgb = BLUE
    _bottom_border(p)


def _bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(0.25)
    p.add_run(text)


def build(resume: dict, path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.left_margin = section.right_margin = Inches(0.8)
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run(resume["name"])
    run.bold = True
    run.font.size = Pt(22)

    first, last = resume["name"].lower().split()
    contact = doc.add_paragraph()
    contact.paragraph_format.space_after = Pt(2)
    run = contact.add_run(f"{resume['location']}  |  {resume['phone']}  |  {first}.{last}@example.com")
    run.font.size = Pt(10)
    run.font.color.rgb = GREY

    _heading(doc, "Summary")
    doc.add_paragraph(resume["summary"]).paragraph_format.space_after = Pt(2)

    _heading(doc, "Experience")
    for role, company, place, dates, bullets in resume["experience"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(role)
        r.bold = True
        p.add_run(f", {company}, {place}")
        r = p.add_run(f"   {dates}")
        r.font.color.rgb = GREY
        r.italic = True
        for bullet in bullets:
            _bullet(doc, bullet)

    _heading(doc, "Education")
    for degree, school, year, note in resume["education"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(degree)
        r.bold = True
        p.add_run(f", {school}, {year}" + (f"  ({note})" if note else ""))

    _heading(doc, "Certifications")
    doc.add_paragraph(" | ".join(resume["certs"])).paragraph_format.space_after = Pt(2)

    _heading(doc, "Skills")
    doc.add_paragraph("  |  ".join(resume["skills"])).paragraph_format.space_after = Pt(2)

    if resume["projects"]:
        _heading(doc, "Projects")
        for project in resume["projects"]:
            _bullet(doc, project)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Fictional test resume for the ASE Candidate Tracker demo. Not a real person.")
    run.font.size = Pt(8)
    run.font.color.rgb = GREY

    doc.save(path)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for resume in RESUMES:
        filename = resume["name"].replace(" ", "_") + "_Resume.docx"
        build(resume, OUT_DIR / filename)
        print("wrote", OUT_DIR / filename)


if __name__ == "__main__":
    main()
