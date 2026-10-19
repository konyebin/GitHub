"""Build samples/FY27Q3_AMER_ASE_Scores_TEST_DATA.xlsx (fictional people) from the blank template."""

import random
from copy import copy
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "samples" / "FY27Q3_AMER_ASE_Scores_TEST_DATA.xlsx"

wb = openpyxl.load_workbook(ROOT / "samples" / "FY27Q3_AMER_ASE_Scores_template.xlsx")
ws = wb.active
first_row = 11
style_row = 11

people = [
    ("USA", "Mr.", "Avery Stone"), ("USA", "Ms.", "Jordan Blake"), ("CAN", "Mx.", "Morgan Ellis"),
    ("USA", "Ms.", "Riley Chen"), ("MEX", "Mr.", "Sam Ortega"), ("BRA", "Ms.", "Taylor Nunes"),
    ("USA", "Mr.", "Casey Brooks"), ("CAN", "Ms.", "Drew Patel"), ("USA", "Ms.", "Quinn Harper"),
    ("COL", "Mr.", "Reese Duarte"), ("USA", "Mx.", "Skyler Moon"), ("CAN", "Mr.", "Parker Lowe"),
    ("USA", "Ms.", "Jamie Fox"), ("ARG", "Mr.", "Alex Rivera"), ("USA", "Ms.", "Hayden Cole"),
    ("MEX", "Ms.", "Rowan Silva"), ("USA", "Mr.", "Emerson Reid"), ("CAN", "Ms.", "Finley Grant"),
    ("BRA", "Mr.", "Logan Costa"), ("USA", "Ms.", "Sage Warren"),
]
strengths = ["Strong storyteller", "Great technical depth", "Calm under pressure", "Needs more depth on AI", "Very coachable"]
buddies = ["P. Morales", "K. Singh", "L. Brandt", "T. Okafor"]
recs = {4: "Yes", 3: "Yes", 2: "Standby", 1: "No"}

random.seed(11)
last_col = ws.max_column
for offset, (country, salutation, name) in enumerate(people):
    r = first_row + offset
    for c in range(1, last_col + 1):
        src = ws.cell(style_row, c)
        dst = ws.cell(r, c)
        if src.has_style:
            dst._style = copy(src._style)
    ws.cell(r, 1, country)
    ws.cell(r, 2, salutation)
    ws.cell(r, 3, name)
    level = random.choice([2, 3, 3, 3, 4, 4])
    for start, rec_col, note_col in ((4, 9, 10), (11, 16, 17), (18, 23, 24)):
        skip_last = name == "Sage Warren" and start == 18  # incomplete candidate: no presentation scores
        scores = []
        for k in range(4):
            if skip_last:
                continue
            value = max(1, min(4, level + random.choice([-1, 0, 0, 1])))
            scores.append(value)
            ws.cell(r, start + k, value)
        if scores:
            ws.cell(r, rec_col, recs[round(sum(scores) / len(scores))])
            ws.cell(r, note_col, random.choice(strengths))
    ws.cell(r, 8, f"=SUM(D{r}:G{r})")
    ws.cell(r, 15, f"=SUM(K{r}:N{r})")
    ws.cell(r, 22, f"=SUM(R{r}:U{r})")
    ws.cell(r, 25, f"=H{r}+O{r}+V{r}")
    ws.cell(r, 26, f"=Y{r}/48")
    if level == 4:
        ws.cell(r, 28, random.choice(buddies))

wb.save(OUT)
print("written", OUT)
