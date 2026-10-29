"""Generate a fully synthetic, already-redacted candidate dataset for developing the analysis.

Nothing here is real. Output goes to hiring-analysis/data/synthetic/ (git-ignored):
  resumes/candidateN.docx      redacted-style resumes ([NAME], [EMAIL], [PHONE] tags; school and city kept)
  screen_review.xlsx           resume screening stage (columns mirror the real screening sheet)
  assessment_center.xlsx       three exercises, 4 competencies each scored 1-4, per-exercise notes, assessor, decision
  _truth.json                  the effects planted below, so we can check the pipeline recovers them

Planted effects (the analysis should find these without being told):
  1. Reviewer R2 scores the Interview exercise about one point per competency higher than the others.
  2. Reviewer R3 writes subjective language ("polished", "great fit", "executive presence") far more often for
     the "career_changer" and "intern" resume paths than for "help_desk", independent of score.
  3. Interview and Role Play ranks agree strongly; Presentation partly diverges (it reflects a separate skill).
  4. Screening Pass/Fail and Status track the TI score, but a few candidates are advanced or stopped against it.
  5. The final Decision follows total score loosely and disagrees for a handful of candidates.
  6. "Buddy" is filled only for hires (outcome leak) and must be excluded from predictors.

Run: python hiring-analysis/synthetic/make_synthetic.py [--n 60] [--seed 7]
Requires: openpyxl, python-docx
"""

from __future__ import annotations

import argparse
import json
import random
from pathlib import Path

import openpyxl
from docx import Document

OUT = Path(__file__).resolve().parents[1] / "data" / "synthetic"

PATHS = ["help_desk", "intern", "career_changer", "bootcamp", "experienced_se"]
CITIES = ["Seattle, WA", "Austin, TX", "Raleigh, NC", "Denver, CO", "Phoenix, AZ", "Chicago, IL", "Atlanta, GA",
          "Boston, MA", "Toronto, ON", "Mexico City, MX", "Sao Paulo, BR", "Bogota, CO"]
SCHOOLS = ["Eastern Carolina University", "Puget Sound State University", "Lakeshore Technical College",
           "University of Georgia", "Northern Plains University", "Metro City College", "Summit Online University"]
EMPLOYERS = ["Harborlight Health", "Evergreen Outfitters", "Brightwater Logistics", "Northfield Bank", "Redwood Retail",
             "Cascade Systems", "Lumen Analytics", "Orchard Insurance"]
SKILLS = ["Active Directory", "Microsoft 365", "ServiceNow", "Basic networking", "Python", "AWS", "Cisco CCNA",
          "Webex Calling", "VoIP", "Salesforce", "SQL", "Customer communication"]
CERTS = ["CompTIA A+", "CompTIA Network+", "Cisco CCNA", "AWS Cloud Practitioner", "Security+"]

INTERVIEW_COMP = ["Growth Mindset, Resilience & Learning Agility", "Executive Presence & Communication",
                  "Collaboration & Teamwork", "Business Acumen & Technical Insights"]
ROLEPLAY_COMP = ["Builds Trust / Trusted Advisor", "Executive Presence & Communication", "Coachability",
                 "Explains Concepts Simply"]
PRESENT_COMP = ["AI Fluency & Orchestration", "Executive Presence & Communication", "Coachability",
                "Goes Deep (Technical Acumen)"]

SUBJECTIVE = ["polished", "great fit", "executive presence", "natural presence", "likeable", "confident delivery",
              "strong gut feel", "would fit the team"]
RUBRIC = ["gave a specific example with measurable results", "walked through the troubleshooting steps in order",
          "tied the answer to the customer's business goal", "asked a clarifying question before answering",
          "described what they would do differently next time", "explained the protocol without jargon",
          "struggled to go deeper than the first layer of the technology", "needed prompting to quantify impact",
          "answered a different question than the one asked", "showed a clear plan for learning the gap",
          "was precise about the trade-offs between two designs"]
FILLER = ["The pace was steady.", "Camera and audio were fine.", "Started a few minutes late.",
          "Took notes throughout the discussion.", "Followed up with a thoughtful closing question."]


def clamp(x: float, lo: int = 1, hi: int = 4) -> int:
    return max(lo, min(hi, int(round(x))))


def make_candidate(i: int, rng: random.Random) -> dict:
    path = rng.choices(PATHS, weights=[3, 2, 2, 2, 1])[0]
    base = {"help_desk": 0.0, "intern": 0.1, "career_changer": -0.1, "bootcamp": -0.1, "experienced_se": 0.5}[path]
    skill = rng.gauss(0, 1) + base
    presence = rng.gauss(0, 1)
    return {"n": i, "path": path, "skill": skill, "presence": presence,
            "city": rng.choice(CITIES), "school": rng.choice(SCHOOLS),
            "grad": f"{rng.choice(['May', 'Dec', 'Aug'])} {rng.choice([2024, 2025, 2026, 2027])}",
            "employers": rng.sample(EMPLOYERS, 2), "skills": rng.sample(SKILLS, 4), "cert": rng.choice(CERTS)}


def score4(latent: float, rng: random.Random, bias: float = 0.0) -> int:
    return clamp(2.5 + 0.7 * latent + bias + rng.gauss(0, 0.6))


def long_note(stage: str, avg: float, cand: dict, reviewer: str, rng: random.Random) -> tuple[str, bool]:
    subj_rate = 0.65 if reviewer == "R3" and cand["path"] in ("career_changer", "intern") else 0.15
    subjective = rng.random() < subj_rate
    parts = [f"{stage} summary."]
    strong = avg >= 3.0
    pool = [r for r in RUBRIC if (r.startswith(("struggled", "needed", "answered a")) != strong)]
    parts += [f"The candidate {x}." for x in rng.sample(pool, k=min(3, len(pool)))]
    if subjective:
        parts.append(f"Overall came across as {rng.choice(SUBJECTIVE)} and {rng.choice(SUBJECTIVE)}.")
    parts += rng.sample(FILLER, 2)
    verdict = "Recommend moving forward." if avg >= 3.0 else ("Borderline; hold for discussion." if avg >= 2.3
                                                              else "Not ready for this role at this time.")
    parts.append(verdict)
    return " ".join(parts), subjective


def write_resume(c: dict, folder: Path) -> None:
    d = Document()
    d.add_paragraph("[NAME]")
    d.add_paragraph(f"{c['city']}  |  [PHONE]  |  [EMAIL]")
    d.add_paragraph("SUMMARY")
    summaries = {
        "help_desk": "Help desk professional who enjoys learning new tools quickly and supporting end users.",
        "intern": "Recent graduate with a summer internship in IT operations and strong interest in customer-facing roles.",
        "career_changer": "Former operations manager moving into pre-sales engineering after self-study and certification.",
        "bootcamp": "Bootcamp graduate with hands-on cloud and Python projects seeking a solutions role.",
        "experienced_se": "Systems engineer with several years of customer-facing design and deployment work.",
    }
    d.add_paragraph(summaries[c["path"]])
    d.add_paragraph("EXPERIENCE")
    for e in c["employers"]:
        d.add_paragraph(f"{e}, {c['city']}")
    d.add_paragraph("EDUCATION")
    d.add_paragraph(f"{c['school']}, {c['grad']}")
    d.add_paragraph("CERTIFICATIONS")
    d.add_paragraph(c["cert"])
    d.add_paragraph("SKILLS")
    d.add_paragraph("  |  ".join(c["skills"]))
    d.save(folder / f"candidate{c['n']}.docx")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--n", type=int, default=60)
    ap.add_argument("--seed", type=int, default=7)
    args = ap.parse_args()
    rng = random.Random(args.seed)
    (OUT / "resumes").mkdir(parents=True, exist_ok=True)

    cands = [make_candidate(i, rng) for i in range(1, args.n + 1)]
    screen_rows, ac_rows, truth = [], [], {"planted": __doc__.split("Planted effects")[1].split("Run:")[0].strip(),
                                             "subjective_notes": {}}

    for c in cands:
        write_resume(c, OUT / "resumes") if rng.random() < 0.85 else None
        ti = clamp(round(21 + 4 * c["skill"] + rng.gauss(0, 3)), 8, 32)
        had_ti = rng.random() < 0.8
        recruiter = rng.choice(["Recruiter 1", "Recruiter 2", "Recruiter 3"])
        passed = had_ti and ti >= 21
        advance = passed if rng.random() > 0.1 else (not passed)
        status = ("Sent Assessment Center Instructions and Candidate Confirmed Availability" if advance and rng.random() < 0.7
                  else "Sent Assessment Center Instructions" if advance else
                  "Not interested in the role" if rng.random() < 0.3 else None)
        c["advance"] = bool(advance and had_ti)
        ti_note, _ = long_note("Technical interview", ti / 8, c, recruiter, rng)
        screen_rows.append([f"candidate{c['n']}", "[NAME]", c["grad"], rng.choice(
            ["Former Intern", "Holds CCNA", "Studying for AWS", "Bachelor Computer Science", "Experience: help desk", "N/A"]),
            c["city"], recruiter, "Sep 2026" if had_ti else "N/A", "Yes" if had_ti else "N/A",
            ti if had_ti else "N/A", None, ("Pass" if passed else "Fail") if had_ti else "N/A",
            ti_note if had_ti else "N/A", ("Yes" if advance else "No") if had_ti else "N/A", status])

    assessors = ["Assessor 1", "Assessor 2", "Assessor 3", "R2", "R3"]
    for c in [x for x in cands if x["advance"]]:
        rows = {}
        total = 0
        for stage, comps, latent in (("Interview", INTERVIEW_COMP, c["skill"] * 0.5 + c["presence"] * 0.5),
                                     ("Role Play", ROLEPLAY_COMP, c["skill"] * 0.4 + c["presence"] * 0.6),
                                     ("Presentation", PRESENT_COMP, c["skill"] * 0.8 + rng.gauss(0, 0.8))):
            reviewer = rng.choice(assessors)
            bias = 1.0 if (stage == "Interview" and reviewer == "R2") else 0.0
            sc = [score4(latent, rng, bias) for _ in comps]
            note, subj = long_note(stage, sum(sc) / 4, c, reviewer, rng)
            truth["subjective_notes"][f"candidate{c['n']}|{stage}"] = {"reviewer": reviewer, "subjective": subj}
            rows[stage] = (sc, reviewer, note, "Yes" if sum(sc) >= 12 else "Standby" if sum(sc) >= 9 else "No")
            total += sum(sc)
        noisy = total + rng.gauss(0, 3)
        decision = "Hire" if noisy >= 36 else "Standby" if noisy >= 28 else "No Hire"
        buddy = "P. Morales" if decision == "Hire" and rng.random() < 0.7 else None
        row = [rng.choice(["USA", "CAN", "MEX", "BRA", "COL"]), rng.choice(["Mr.", "Ms.", "Mx."]), f"candidate{c['n']}"]
        for stage in ("Interview", "Role Play", "Presentation"):
            sc, reviewer, note, rec = rows[stage]
            row += sc + [sum(sc), rec, note, reviewer]
        row += [total, round(total / 48, 3), decision, buddy]
        ac_rows.append(row)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Screen"
    ws.append(["Person ID", "Full name", "Graduation Date or Expected Graduation Date", "Skills and Technologies",
               "Home cities", "Recruiter", "TI Date", "Confirmed Yes/No", "Score (Passing Score 21)",
               "Additional Comments", "Pass/Fail", "Notes (Summary of the TI call)",
               "Recommended to be Invited to Insight Day Yes/No", "Status"])
    for r in screen_rows:
        ws.append(r)
    wb.save(OUT / "screen_review.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Scores"
    head = ["Country", "Salutation", "Candidate"]
    for stage, comps in (("Interview", INTERVIEW_COMP), ("Role Play", ROLEPLAY_COMP), ("Presentation", PRESENT_COMP)):
        head += [f"{stage}: {c}" for c in comps] + [f"TOTAL - {stage}", f"Recommendation - {stage}",
                                                     f"Comments - {stage}", f"Assessor - {stage}"]
    head += ["Total Score", "Overall %", "Decision", "Buddy"]
    ws.append(head)
    for r in ac_rows:
        ws.append(r)
    wb.save(OUT / "assessment_center.xlsx")

    (OUT / "_truth.json").write_text(json.dumps(truth, indent=2))
    print(f"wrote {args.n} candidates ({len(ac_rows)} in assessment center) to {OUT}")


if __name__ == "__main__":
    main()
